"""Форматные документы: pdf, docx, xlsx, pptx, csv, html, xml, fb2, epub.

Чистая логика без Qt — тестируется отдельно. Тяжёлые библиотеки импортируются
лениво, чтобы не замедлять старт приложения.
Лицензии зависимостей: pypdfium2 (Apache-2.0, движок Pdfium), pypdf (BSD),
python-docx (MIT), mammoth (BSD-2), openpyxl (MIT), python-pptx (MIT) —
свободное ПО без компонентов Adobe и без GPL/AGPL.
"""

from __future__ import annotations

import csv
import io
import os
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from html import escape

DOC_KINDS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "doc",
    ".rtf": "rtf",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".xls": "xls",
    ".pptx": "pptx",
    ".csv": "csv",
    ".tsv": "csv",
    ".html": "html",
    ".htm": "html",
    ".xhtml": "html",
    ".xml": "xml",
    ".fb2": "fb2",
    ".epub": "epub",
}

XLSX_MAX_ROWS = 1_000_000
XLSX_MAX_COLS = 256
GRID_BATCH = 5000
CSV_DELIMS = ",;\t|"


def document_kind(path: str) -> str | None:
    return DOC_KINDS.get(os.path.splitext(path)[1].lower())


def _atomic_replace_write(out_path: str, writer) -> None:
    fd, tmp = tempfile.mkstemp(prefix=".sc_pdf_", suffix=".pdf",
                               dir=os.path.dirname(out_path) or ".")
    os.close(fd)
    try:
        with open(tmp, "wb") as f:
            writer.write(f)
        os.replace(tmp, out_path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


# ---------------------------------------------------------------- pdf

def pdf_page_count(path: str) -> int:
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(path) as pdf:
        return len(pdf)


def pdf_render(path: str, index: int, scale: float = 2.0,
               draw_annots: bool = True) -> tuple[bytes, int, int, int]:
    """Отрисовать страницу: (BGR-данные, ширина, высота, stride)."""
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(path) as pdf:
        page = pdf[index]
        bitmap = page.render(scale=scale, draw_annots=draw_annots)
        return bytes(bitmap.buffer), bitmap.width, bitmap.height, bitmap.stride


def pdf_page_text(path: str, index: int) -> str:
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(path) as pdf:
        page = pdf[index]
        textpage = page.get_textpage()
        return textpage.get_text_range()


def pdf_rotate_pages(path: str, indices: list[int], delta: int) -> None:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(path)
    writer = PdfWriter()
    wanted = set(indices)
    for i, page in enumerate(reader.pages):
        page_obj = page
        if i in wanted:
            page_obj = page.rotate(delta)
        writer.add_page(page_obj)
    _atomic_replace_write(path, writer)


def pdf_delete_pages(path: str, indices: list[int]) -> None:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(path)
    writer = PdfWriter()
    skip = set(indices)
    for i, page in enumerate(reader.pages):
        if i not in skip:
            writer.add_page(page)
    _atomic_replace_write(path, writer)


def pdf_export_pages(path: str, indices: list[int], out_path: str) -> None:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(path)
    writer = PdfWriter()
    for i in indices:
        writer.add_page(reader.pages[i])
    with open(out_path, "wb") as f:
        writer.write(f)


def parse_ranges(text: str, maximum: int) -> list[int]:
    """«1-3,5» → [0, 1, 2, 4]; ValueError при мусоре."""
    result: list[int] = []
    for part in text.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, _, b = part.partition("-")
            start, end = int(a) - 1, int(b) - 1
            if start < 0 or end < start:
                raise ValueError(part)
            result.extend(range(start, end + 1))
        else:
            n = int(part) - 1
            if n < 0:
                raise ValueError(part)
            result.append(n)
    if not result:
        raise ValueError("пусто")
    return sorted({n for n in result if 0 <= n < maximum})


# -- аннотации pdf ---------------------------------------------------------

ANNOT_KINDS = ("highlight", "text", "freetext")


def pdf_page_size(path: str, index: int) -> tuple[float, float]:
    """(ширина, высота) страницы в пунктах PDF."""
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(path) as pdf:
        return tuple(pdf[index].get_size())


def pdf_annot_list(path: str) -> list[dict]:
    """Список аннотаций: [{page, index, subtype, rect, contents}].
    rect — (x0, y0, x1, y1) в координатах PDF (начало — левый нижний угол)."""
    from pypdf import PdfReader

    out = []
    reader = PdfReader(path)
    for page_num, page in enumerate(reader.pages):
        annots = page.get("/Annots")
        if not annots:
            continue
        for i, ref in enumerate(annots):
            try:
                obj = ref.get_object()
                rect = tuple(round(float(v), 2) for v in obj.get("/Rect",
                                                             (0, 0, 0, 0)))
                contents = str(obj.get("/Contents", "") or "")
                subtype = str(obj.get("/Subtype", "")).lstrip("/")
                out.append({"page": page_num, "index": i, "subtype": subtype,
                            "rect": rect, "contents": contents})
            except Exception:
                continue
    return out


def pdf_add_annotation(path: str, page: int, kind: str,
                       rect: tuple[float, float, float, float],
                       contents: str = "") -> None:
    """Добавить аннотацию (highlight/text/freetext) и перезаписать файл.
    rect в координатах PDF; подсветка получает QuadPoints из прямоугольника."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.annotations import FreeText, Highlight, Text
    from pypdf.generic import ArrayObject, FloatObject

    reader = PdfReader(path)
    writer = PdfWriter()
    writer.append(reader)
    if kind == "highlight":
        x0, y0, x1, y1 = rect
        quads = ArrayObject(
            [FloatObject(v) for v in (x0, y1, x1, y1, x0, y0, x1, y0)])
        annot = Highlight(rect=rect, quad_points=quads,
                          highlight_color="ffe066", printing=True)
    elif kind == "text":
        annot = Text(rect=rect, text=contents)
    elif kind == "freetext":
        annot = FreeText(text=contents, rect=rect, font_size="10pt",
                         font_color="202020", border_color="a6784f",
                         background_color="fff8e1")
    else:
        raise ValueError(f"неизвестный тип аннотации: {kind}")
    writer.add_annotation(page_number=page, annotation=annot)
    _atomic_replace_write(path, writer)


def pdf_delete_annotation(path: str, page: int, index: int) -> None:
    """Удалить аннотацию по её индексу в /Annots страницы."""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, NameObject

    reader = PdfReader(path)
    writer = PdfWriter()
    writer.append(reader)
    page_obj = writer.pages[page]
    annots = page_obj.get("/Annots")
    if annots:
        rest = [a for i, a in enumerate(annots) if i != index]
        if rest:
            page_obj[NameObject("/Annots")] = ArrayObject(rest)
        else:
            del page_obj[NameObject("/Annots")]
    _atomic_replace_write(path, writer)


# ---------------------------------------------------------------- docx

def docx_paragraphs(path: str) -> list[str]:
    from docx import Document

    return [p.text for p in Document(path).paragraphs]


def docx_save_paragraphs(path: str, lines: list[str]) -> None:
    """Правка по абзацам: текст абзаца заменяется, его стиль сохраняется
    (встроенное форматирование внутри изменённого абзаца теряется)."""
    from docx import Document

    document = Document(path)
    paragraphs = document.paragraphs
    for i, paragraph in enumerate(paragraphs):
        new_text = lines[i] if i < len(lines) else None
        if new_text is None:
            paragraph._element.getparent().remove(paragraph._element)
        elif paragraph.text != new_text:
            paragraph.text = new_text
    if len(lines) > len(paragraphs):
        for line in lines[len(paragraphs):]:
            document.add_paragraph(line)
    fd, tmp = tempfile.mkstemp(prefix=".sc_docx_", suffix=".docx",
                               dir=os.path.dirname(path) or ".")
    os.close(fd)
    try:
        document.save(tmp)
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def docx_to_html(path: str) -> str:
    import mammoth

    with open(path, "rb") as f:
        return mammoth.convert_to_html(f).value


# ---------------------------------------------------------------- xlsx

def xlsx_sheet_names(path: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def xlsx_sheets(path: str) -> list[tuple[str, list[list[str]]]]:
    """Все листы целиком (строки-списки строк). Для больших файлов лучше
    xlsx_rows_iter + фоновая загрузка."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        for name in wb.sheetnames:
            rows = [list(r) for r in xlsx_rows_iter(path, name)]
            sheets.append((name, rows))
        return sheets
    finally:
        wb.close()


def xlsx_rows_iter(path: str, sheet_name: str):
    """Итератор строк листа (списки строк); для фоновой загрузки большими
    листами — модель получает данные порциями, GUI не блокируется."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row[:XLSX_MAX_COLS]]
            while values and values[-1] == "":
                values.pop()  # хвостовые пустые ячейки не храним
            yield values
    finally:
        wb.close()


def xlsx_sheet_raw(path: str, sheet_name: str):
    """Сырое содержимое листа для правки и пересчёта формул.

    Возвращает (rows_raw, formulas, cached):
      rows_raw — строки сырых ячеек (формулы строками '=…', значения строками);
      formulas — {(строка, колонка): '=…'} координаты формул;
      cached   — {(строка, колонка): прежнее кэшированное значение} — seed
                 для движка пересчёта (formulas.evaluate_sheet).
    """
    from openpyxl import load_workbook

    rows_raw: list[list[str]] = []
    formulas: dict[tuple, str] = {}
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        ws = wb[sheet_name]
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            values = []
            for c, v in enumerate(row[:XLSX_MAX_COLS]):
                if isinstance(v, str) and v.startswith("="):
                    formulas[(r, c)] = v
                values.append("" if v is None else str(v))
            while values and values[-1] == "":
                values.pop()
            rows_raw.append(values)
    finally:
        wb.close()
    cached: dict[tuple, object] = {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            for c, v in enumerate(row[:XLSX_MAX_COLS]):
                if v is not None:
                    cached[(r, c)] = v
    finally:
        wb.close()
    return rows_raw, formulas, cached


def xlsx_save_cells(path: str, edited: dict[str, list[list[str]]],
                    originals: dict[str, list[list[str]]]) -> None:
    """Записать листы, изменённые в редакторе сетки.

    edited — {имя листа: строки сырых ячеек}; originals — их состояние при
    загрузке: неизменившиеся ячейки не переписываются (типы — даты, стили —
    сохраняются), меняются только правленые.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for name, rows in edited.items():
        orig = originals.get(name) or []
        ws = wb[name]
        width = max(len(r) for r in rows) if rows else 0
        for r, row in enumerate(rows):
            old = orig[r] if r < len(orig) else []
            for c in range(max(len(row), len(old))):
                raw = row[c] if c < len(row) else ""
                if c < len(old) and old[c] == raw:
                    continue  # не менялось — тип и стиль сохраняются
                ws.cell(row=r + 1, column=c + 1, value=_xlsx_typed(raw))
    fd, tmp = tempfile.mkstemp(prefix=".sc_xlsx_", suffix=".xlsx",
                               dir=os.path.dirname(path) or ".")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _xlsx_typed(raw: str):
    """Сырая строка сетки → значение для openpyxl (формула/число/логич./текст)."""
    s = raw.strip()
    if s.startswith("="):
        return raw
    if s.upper() == "TRUE":
        return True
    if s.upper() == "FALSE":
        return False
    num = s.replace(" ", "").replace(",", ".")
    try:
        f = float(num)
        return int(f) if f.is_integer() and abs(f) < 1e15 else f
    except ValueError:
        return raw


def _xlsx_sheet_xml_map(zf: zipfile.ZipFile) -> dict[str, str]:
    """Имя листа → член архива с его XML (через workbook.xml + rels)."""
    WB_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    REL_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    PKG_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {}
    for rel in rels.findall(f"{{{PKG_NS}}}Relationship"):
        if rel.get("Type", "").endswith("/worksheet"):
            target = rel.get("Target", "")
            if target.startswith("/"):
                target = target[1:]
            elif not target.startswith("xl/"):
                target = "xl/" + target
            targets[rel.get("Id")] = target
    out = {}
    for sheet in workbook.iter(f"{{{WB_NS}}}sheet"):
        rid = sheet.get(REL_ID)
        if rid in targets:
            out[sheet.get("name")] = targets[rid]
    return out


def xlsx_inject_cached(path: str, cache: dict[str, dict[tuple, object]]) -> None:
    """Вписать кэшированные значения формул в XML листов (после openpyxl.save:
    openpyxl не пишет <v> у формул, и читатели без пересчёта показали бы пусто).
    cache — {имя листа: {(строка, колонка): значение}}; координаты 0-based.
    """
    if not cache:
        return
    sheet_map = None
    fd, tmp = tempfile.mkstemp(prefix=".sc_xlsx_v_", suffix=".xlsx",
                               dir=os.path.dirname(path) or ".")
    os.close(fd)
    try:
        with zipfile.ZipFile(path) as zin:
            sheet_map = _xlsx_sheet_xml_map(zin)
            wanted = {sheet_map[n]: n for n in cache if n in sheet_map}
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
                for member in zin.namelist():
                    data = zin.read(member)
                    if member in wanted:
                        data = _xlsx_patch_sheet_xml(
                            data, cache[wanted[member]])
                    zout.writestr(member, data)
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _xlsx_patch_sheet_xml(data: bytes, values: dict[tuple, object]) -> bytes:
    """Вставить <v> в ячейки с <f>: числа/логические/строки/ошибки Excel."""
    NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", NS)
    root = ET.fromstring(data)
    changed = False
    for cell in root.iter(f"{{{NS}}}c"):
        formula = cell.find(f"{{{NS}}}f")
        if formula is None:
            continue
        ref = cell.get("r", "")
        col_letters = "".join(ch for ch in ref if ch.isalpha())
        row_digits = "".join(ch for ch in ref if ch.isdigit())
        if not col_letters or not row_digits:
            continue
        col = 0
        for ch in col_letters.upper():
            col = col * 26 + (ord(ch) - 64)
        coord = (int(row_digits) - 1, col - 1)
        if coord not in values:
            continue
        v = values[coord]
        for old in cell.findall(f"{{{NS}}}v"):
            cell.remove(old)
        ve = ET.SubElement(cell, f"{{{NS}}}v")
        from sphaera_commander.formulas import ExcelError
        from sphaera_commander.formulas import fmt_number_for_xml

        if isinstance(v, ExcelError):
            cell.set("t", "e")
            ve.text = v.code
        elif isinstance(v, bool):
            cell.set("t", "b")
            ve.text = "1" if v else "0"
        elif isinstance(v, (int, float)):
            if "t" in cell.attrib:
                del cell.attrib["t"]
            ve.text = fmt_number_for_xml(v)
        else:
            cell.set("t", "str")
            ve.text = str(v)
        changed = True
    if not changed:
        return data
    return ET.tostring(root, encoding="UTF-8", xml_declaration=True)


def csv_rows_iter(text: str, delimiter: str):
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    count = 0
    for row in reader:
        if count >= XLSX_MAX_ROWS:
            return
        count += 1
        yield row


def csv_sniff_delimiter(text: str) -> str:
    head = "\n".join(text.splitlines()[:10])
    best, best_score = ",", -1
    for d in CSV_DELIMS:
        score = head.count(d)
        if score > best_score:
            best, best_score = d, score
    return best if best_score > 0 else ","


# ---------------------------------------------------------------- pptx

def pptx_slides(path: str) -> list[tuple[int, str]]:
    from pptx import Presentation

    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if shape == slide.shapes.title:
                title = text
            parts.append(escape(text).replace("\n", "<br>"))
        body = "<hr>".join(parts) if parts else "<i>(пусто)</i>"
        slides.append((i, f"<b>{escape(title)}</b><hr>{body}"))
    return slides


# Цвета темы Office по умолчанию — для фигур без явного RGB
THEME_RGB = {
    "dk1": 0x000000, "lt1": 0xFFFFFF, "dk2": 0x44546A, "lt2": 0xE7E6E6,
    "accent1": 0x4472C4, "accent2": 0xED7D31, "accent3": 0xA5A5A5,
    "accent4": 0xFFC000, "accent5": 0x5B9BD5, "accent6": 0x70AD47,
    "hlink": 0x0563C1, "folHlink": 0x954F72, "phClr": 0x000000,
    "tx1": 0x000000, "bg1": 0xFFFFFF, "tx2": 0x44546A, "bg2": 0xE7E6E6,
}


def _color_rgb(color) -> int | None:
    """RGB цвета python-pptx; темы маппятся на палитру Office."""
    try:
        from pptx.enum.dml import MSO_COLOR_TYPE

        if color.type == MSO_COLOR_TYPE.RGB:
            return int(str(color.rgb), 16)
        if color.type == MSO_COLOR_TYPE.SCHEME:
            return THEME_RGB.get(str(color.theme_color).rsplit("(", 1)[0]
                                 .strip().lower(), None)
    except Exception:
        pass
    return None


def _fill_rgb(fill) -> int | None:
    try:
        from pptx.enum.dml import MSO_FILL

        if fill.type == MSO_FILL.SOLID:
            return _color_rgb(fill.fore_color)
    except Exception:
        pass
    return None


def _para_html(paragraph) -> str:
    from pptx.util import Pt

    runs = []
    for run in paragraph.runs:
        font = run.font
        size = font.size.pt if font.size else None
        bold = " font-weight:bold;" if font.bold else ""
        italic = " font-style:italic;" if font.italic else ""
        color = ""
        rgb = _color_rgb(font.color)
        if rgb is not None:
            color = f" color:#{rgb:06x};"
        family = f" font-family:'{font.name}';" if font.name else ""
        size_style = f" font-size:{size:.0f}pt;" if size else ""
        text = escape(run.text).replace("\v", "<br>")
        runs.append(f"<span style=\"{size_style}{bold}{italic}{color}{family}\">{text}</span>")
    align = {1: "center", 2: "right", 3: "justify"}.get(
        int(paragraph.alignment) if paragraph.alignment is not None else 0, "left")
    level = min(int(paragraph.level or 0), 5)
    html = "".join(runs) or "&nbsp;"
    return (f"<p align=\"{align}\" style=\"margin:2 0 2 "
            f"{level * 18}px;\">{html}</p>")


def _shape_box(shape, images_dir: str | None, prefix: str) -> dict:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    box = {"x": int(shape.left or 0), "y": int(shape.top or 0),
           "w": int(shape.width or 0), "h": int(shape.height or 0),
           "kind": "rect", "fill": None, "line": None, "img": None,
           "html": "", "anchor": "top", "table": None}
    try:
        from pptx.enum.shapes import MSO_SHAPE

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            box["kind"] = "picture"
            if images_dir is not None:
                image = shape.image
                img_path = os.path.join(images_dir, f"{prefix}.{image.ext}")
                with open(img_path, "wb") as f:
                    f.write(image.blob)
                box["img"] = img_path
            return box
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            name = str(getattr(shape, "auto_shape_type", "")).upper()
            if "OVAL" in name or "ELLIPSE" in name or "CIRCLE" in name:
                box["kind"] = "ellipse"
    except Exception:
        pass
    box["fill"] = _fill_rgb(shape.fill)
    try:
        line_rgb = _fill_rgb(shape.line.fill)
        box["line"] = line_rgb
    except Exception:
        box["line"] = None
    try:
        if shape.has_table:
            box["kind"] = "table"
            box["table"] = [[cell.text for cell in row.cells]
                            for row in shape.table.rows]
            return box
    except Exception:
        pass
    try:
        if shape.has_text_frame:
            from pptx.enum.text import MSO_VERTICAL_ANCHOR

            tf = shape.text_frame
            anchor = {MSO_VERTICAL_ANCHOR.MIDDLE: "middle",
                      MSO_VERTICAL_ANCHOR.BOTTOM: "bottom"}.get(
                          tf.vertical_anchor, "top")
            box["anchor"] = anchor
            box["html"] = "".join(_para_html(p) for p in tf.paragraphs)
    except Exception:
        pass
    return box


def pptx_slides_rich(path: str, images_dir: str | None = None) -> dict:
    """Геометрия слайдов для собственного рендера (без LibreOffice).

    Поддерживается: сплошной фон, прямоугольники/овалы с заливкой и рамкой,
    картинки, таблицы (упрощённо), текст с кеглем/цветом/выравниванием
    и переносом. Градиенты, тени, SmartArt, диаграммы упрощаются.
    """
    from pptx import Presentation

    prs = Presentation(path)
    deck = {"width": int(prs.slide_width), "height": int(prs.slide_height),
            "slides": []}
    for i, slide in enumerate(prs.slides):
        bg = None
        try:
            bg = _fill_rgb(slide.background.fill)
        except Exception:
            bg = None
        deck["slides"].append({
            "bg": bg,
            "shapes": [_shape_box(shape, images_dir, f"slide{i + 1}_{j}")
                       for j, shape in enumerate(slide.shapes)],
        })
    return deck


# ---------------------------------------------------------------- csv

def csv_grid(text: str) -> tuple[str, list[list[str]]]:
    """Определить разделитель и разобать; (разделитель, строки)."""
    delim = csv_sniff_delimiter(text)
    rows = list(csv_rows_iter(text, delim))
    return delim, rows


# ---------------------------------------------------------------- fb2

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def fb2_html(path: str) -> tuple[str, str]:
    """(название книги, html основных разделов)."""
    tree = ET.parse(path)
    root = tree.getroot()
    title = ""
    for el in root.iter():
        if _local(el.tag) == "book-title" and el.text:
            title = el.text.strip()
            break
    body = next((el for el in root.iter() if _local(el.tag) == "body"), None)
    parts: list[str] = []

    def walk(node) -> None:
        for child in node:
            name = _local(child.tag)
            if name == "title":
                text = " ".join(t.strip() for t in child.itertext() if t.strip())
                if text:
                    parts.append(f"<h3>{escape(text)}</h3>")
            elif name == "p":
                inner = "".join(_fb2_inline(c) for c in child)
                parts.append(f"<p>{escape(child.text or '')}{inner}</p>")
            elif name in ("poem", "epigraph"):
                text = escape(" ".join(t.strip() for t in child.itertext() if t.strip()))
                parts.append(f"<p><i>{text}</i></p>")
            elif name == "section":
                walk(child)
            elif name == "empty-line":
                parts.append("<br>")

    if body is not None:
        walk(body)
    return title, "\n".join(parts) or "<i>(пусто)</i>"


def _fb2_inline(node) -> str:
    name = _local(node.tag)
    text = escape(node.text or "")
    tail = escape(node.tail or "")
    inner = "".join(_fb2_inline(c) for c in node)
    if name in ("emphasis", "style", "strong"):
        return f"<i>{text}{inner}</i>{tail}"
    if name == "a":
        return f"<u>{text}{inner}</u>{tail}"
    if name == "subtitle":
        return f"<b>{text}{inner}</b>{tail}"
    return f"{text}{inner}{tail}"


# ---------------------------------------------------------------- epub

def epub_chapters(path: str, extract_to: str | None = None) -> list[tuple[str, str, str | None]]:
    """Главы по spine: (заголовок, исходный xhtml, базовый каталог).

    Если extract_to задан — архив распаковывается туда и base — реальный
    каталог главы: относительные <img> и ссылки разрешаются при рендере
    (QTextBrowser с baseUrl). Иначе base=None (только текст/разметка).
    """
    with zipfile.ZipFile(path) as zf:
        container = ET.fromstring(zf.read("META-INF/container.xml"))
        rootfile = next(el for el in container.iter()
                        if _local(el.tag) == "rootfile")
        opf_path = rootfile.attrib["full-path"]
        base = os.path.dirname(opf_path)
        opf = ET.fromstring(zf.read(opf_path))

        manifest: dict[str, tuple[str, str]] = {}
        for item in opf.iter():
            if _local(item.tag) == "item":
                manifest[item.attrib["id"]] = (item.attrib["href"],
                                               item.attrib.get("media-type", ""))
        spine = [el.attrib["idref"] for el in opf.iter()
                 if _local(el.tag) == "itemref"]

        if extract_to is not None:
            zf.extractall(extract_to)

        def real_base(href: str) -> str | None:
            if extract_to is None:
                return None
            full = os.path.normpath(os.path.join(extract_to, base, os.path.dirname(href)))
            return full if os.path.isdir(full) else extract_to

        chapters = []
        for idref in spine:
            href, media = manifest.get(idref, ("", ""))
            if media and "html" not in media:
                continue
            full = f"{base}/{href}" if base else href
            try:
                raw = zf.read(full)
            except KeyError:
                continue
            xhtml = raw.decode("utf-8", errors="replace")
            title = _html_title(xhtml) or f"Глава {len(chapters) + 1}"
            chapters.append((title, xhtml, real_base(href)))
        return chapters


def _html_title(xhtml: str) -> str:
    match = re.search(r"<title[^>]*>(.*?)</title>", xhtml, re.S | re.I)
    if match:
        return " ".join(match.group(1).split())
    match = re.search(r"<h[1-3][^>]*>(.*?)</h[1-3]>", xhtml, re.S | re.I)
    if match:
        return " ".join(re.sub(r"<[^>]+>", "", match.group(1)).split())
    return ""


# ---------------------------------------------------------------- xml

def xml_tree(path: str) -> tuple[str, tuple]:
    """(имя корня, узел) где узел = (тег, текст, [дети])."""
    root = ET.parse(path).getroot()

    def build(el) -> tuple:
        return (_local(el.tag), " ".join((el.text or "").split()),
                [build(child) for child in el])

    return _local(root.tag), build(root)
